#!/usr/bin/env python3
"""
Create GitHub issues from a LOT_x/*.xlsx article spreadsheet.

Invoked by .github/workflows/create-lot-issues.yml whenever a push's commit
message contains a LOT_<n> reference (e.g. "LOT_1"). Reads LOT_<n>/*.xlsx,
maps each row onto the "Blog Article" issue template, and opens one issue
per row -- skipping rows that already have a matching issue (dedup via a
hidden marker embedded in the issue body).

If a row's HTML content cell is empty, falls back to the matching .html
file inside LOT_<n>/*.zip (matched by slugified title).
"""

import os
import re
import sys
import glob
import unicodedata
import zipfile
import argparse

import requests
import openpyxl

API_ROOT = "https://api.github.com"

# Header names expected in the spreadsheet (must match exactly).
REQUIRED_HEADERS = {
    "title": "Titre de l'article",
    "description": "Description SEO (meta description)",
    "category": "Category",
    "content": "Contenu HTML de l'article",
}
ID_HEADER = "N°"


def slugify(value: str) -> str:
    """ASCII-safe, lowercase, hyphenated slug (used for label values)."""
    value = str(value or "")
    # drop emoji / symbol characters (e.g. flag emojis) before transliterating
    value = "".join(ch for ch in value if not unicodedata.category(ch).startswith("So"))
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", "-", value)
    return value.strip("-") or "uncategorized"


def detect_language(html: str, default: str = "fr") -> str:
    match = re.search(r'lang=["\']?([a-z]{2})["\']?', html or "", re.IGNORECASE)
    return match.group(1).lower() if match else default


def find_sheet(wb):
    """Return the worksheet whose header row has all the required columns."""
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        if all(h in headers for h in REQUIRED_HEADERS.values()):
            return ws, headers
    raise RuntimeError(
        "No sheet found with the expected columns: " + ", ".join(REQUIRED_HEADERS.values())
    )


def load_fallback_html(zip_path, title):
    """If a row's HTML cell is empty, try to recover it from the zip by slug match."""
    if not zip_path or not os.path.exists(zip_path):
        return None
    target_slug = slugify(title)
    with zipfile.ZipFile(zip_path) as zf:
        for name in zf.namelist():
            if name.endswith(".html"):
                base = os.path.splitext(os.path.basename(name))[0]
                if slugify(base) == target_slug:
                    return zf.read(name).decode("utf-8")
    return None


def build_marker(lot_id: str, row_id) -> str:
    # Alphanumeric-only so GitHub's issue-body full-text search treats it as a single token.
    clean_lot = re.sub(r"[^A-Za-z0-9]", "", lot_id).upper()
    return f"LOTROWMARKER{clean_lot}ROW{int(row_id):04d}"


def build_body(description, category, language, content_html, marker: str) -> str:
    # Mirrors the markdown GitHub renders for a submitted issue form, so any
    # downstream parsing built against the manual-submission format still works.
    return (
        "### Description\n\n"
        f"{description}\n\n"
        "### Category\n\n"
        f"{category}\n\n"
        "### Language\n\n"
        f"{language}\n\n"
        "### Featured\n\n"
        "- [ ] Feature this article\n\n"
        "### Content\n\n"
        f"{content_html}\n\n"
        f"<!-- {marker} -->\n"
    )


def get_existing_markers(session, owner, repo):
    """Fetch all issues (any state) once and collect embedded dedup markers."""
    markers = set()
    page = 1
    while True:
        resp = session.get(
            f"{API_ROOT}/repos/{owner}/{repo}/issues",
            params={"state": "all", "per_page": 100, "page": page},
        )
        resp.raise_for_status()
        batch = resp.json()
        if not batch:
            break
        for issue in batch:
            if "pull_request" in issue:  # the issues endpoint also lists PRs
                continue
            body = issue.get("body") or ""
            markers.update(re.findall(r"LOTROWMARKER[A-Z0-9]+", body))
        page += 1
    return markers


def create_issue(session, owner, repo, title, body, labels):
    resp = session.post(
        f"{API_ROOT}/repos/{owner}/{repo}/issues",
        json={"title": title, "body": body, "labels": labels},
    )
    resp.raise_for_status()
    return resp.json()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--lot", required=True, help="LOT folder name, e.g. LOT_1")
    args = parser.parse_args()

    token = os.environ["GITHUB_TOKEN"]
    owner, repo = os.environ["GITHUB_REPOSITORY"].split("/")

    session = requests.Session()
    session.headers.update(
        {
            "Authorization": f"Bearer {token}",
            "Accept": "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
        }
    )

    lot_id = args.lot

    xlsx_matches = sorted(glob.glob(os.path.join(lot_id, "*.xlsx")))
    if not xlsx_matches:
        print(f"::warning::No .xlsx file found under {lot_id}/ -- nothing to do.")
        return
    xlsx_path = xlsx_matches[0]
    if len(xlsx_matches) > 1:
        print(f"::warning::Multiple .xlsx files under {lot_id}/, using {xlsx_path}")

    zip_matches = sorted(glob.glob(os.path.join(lot_id, "*.zip")))
    zip_path = zip_matches[0] if zip_matches else None

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws, headers = find_sheet(wb)
    col_index = {name: idx for idx, name in enumerate(headers)}
    id_col = col_index.get(ID_HEADER)

    print("Fetching existing issues to build the dedup index...")
    existing_markers = get_existing_markers(session, owner, repo)
    print(f"Found {len(existing_markers)} existing LOT markers across all issues.")

    created = skipped = errors = 0

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue

        title = row[col_index[REQUIRED_HEADERS["title"]]]
        description = row[col_index[REQUIRED_HEADERS["description"]]]
        category = row[col_index[REQUIRED_HEADERS["category"]]]
        content = row[col_index[REQUIRED_HEADERS["content"]]]
        row_id = row[id_col] if id_col is not None and row[id_col] is not None else row_number

        if not title:
            continue

        if not content:
            content = load_fallback_html(zip_path, title)
        if not content:
            print(f"::error::Row {row_id} ('{title}') has no HTML content anywhere -- skipping.")
            errors += 1
            continue

        language = detect_language(content)
        marker = build_marker(lot_id, row_id)

        if marker in existing_markers:
            print(f"Skipping row {row_id} ('{title}') -- an issue already exists for it.")
            skipped += 1
            continue

        issue_title = f"[Article]: {title}"
        body = build_body(description or "", category or "", language, content, marker)
        labels = ["blog", f"category:{slugify(category)}", f"lang:{language}"]

        try:
            issue = create_issue(session, owner, repo, issue_title, body, labels)
            print(f"Created issue #{issue['number']} for row {row_id}: {title}")
            created += 1
        except requests.HTTPError as exc:
            print(f"::error::Failed to create issue for row {row_id} ('{title}'): {exc}")
            errors += 1

    print(f"\nDone. Created={created} Skipped={skipped} Errors={errors}")
    if errors:
        sys.exit(1)


if __name__ == "__main__":
    main()
